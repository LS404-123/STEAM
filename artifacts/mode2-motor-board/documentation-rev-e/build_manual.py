"""從保留的 System Design 範本製作 E 版使用說明，不修改 KiCad 設計。"""
from copy import deepcopy
from pathlib import Path
from zipfile import ZipFile, ZIP_DEFLATED
from hashlib import sha256
import json
import shutil
from lxml import etree
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.text.paragraph import Paragraph
from docx.table import Table
from docx.enum.text import WD_ALIGN_PARAGRAPH

ROOT = Path(__file__).parent
BOARD = ROOT.parent/'kicad-rev-e'
REF = Path('C:/Users/LS404/.codex/plugins/cache/openai-curated-remote/openai-templates/0.1.1/skills/artifact-template-system-design/assets/reference.docx')
OUT = ROOT/'模式二控制板E版使用與原理說明.docx'
template = Document(REF)
doc = Document(REF)
body = doc._element.body
for el in list(body):
    if el.tag != qn('w:sectPr'):
        body.remove(el)

def font(run, size=None, bold=None, color=None):
    if size: run.font.size = Pt(size)
    if bold is not None: run.bold = bold
    if color: run.font.color.rgb = RGBColor.from_string(color)
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'), 'Microsoft JhengHei')
    return run

def p(text='', source=22, size=11, bold=False):
    el = deepcopy(template.paragraphs[source]._p)
    body.insert(len(body)-1, el)
    par = Paragraph(el, doc._body)
    par.clear()
    par.paragraph_format.keep_with_next = False
    par.paragraph_format.widow_control = True
    font(par.add_run(text), size, bold, '233447')
    return par

def heading(text, newpage=True):
    par = p(text, source=21, size=13.5, bold=True)
    par.paragraph_format.page_break_before = newpage
    par.paragraph_format.keep_with_next = True
    par.paragraph_format.space_before = Pt(0)
    for r in par.runs: r.font.color.rgb = RGBColor(0,0,0)

def sub(text):
    par = p(text, source=35, size=11, bold=True)
    par.paragraph_format.space_before = Pt(9)
    par.paragraph_format.keep_with_next = True
    for r in par.runs: r.font.color.rgb = RGBColor(0,0,0)

def picture(path, width, caption=None):
    par = p('', source=32)
    par.paragraph_format.space_after = Pt(4)
    par.alignment = WD_ALIGN_PARAGRAPH.CENTER
    shape = par.add_run().add_picture(str(path), width=Inches(width))
    shape._inline.docPr.set('descr', caption or Path(path).stem)
    if caption:
        par.paragraph_format.keep_with_next = True
        cap = p(caption, source=33, size=9)
        cap.paragraph_format.space_after = Pt(9)
    return par

def table(headers, rows, widths):
    src = template.tables[5]
    el = deepcopy(src._tbl)
    body.insert(len(body)-1, el)
    t = Table(el, doc._body)
    for row in list(el.findall(qn('w:tr'))): el.remove(row)
    grid = el.find(qn('w:tblGrid'))
    for child in list(grid): grid.remove(child)
    for width in widths:
        col = OxmlElement('w:gridCol'); col.set(qn('w:w'), str(round(width*1440))); grid.append(col)
    el.find(qn('w:tblPr')).find(qn('w:tblW')).set(qn('w:w'), str(round(sum(widths)*1440)))
    for i, values in enumerate([headers]+rows):
        row = deepcopy(src.rows[0 if i == 0 else 1]._tr)
        for cell in list(row.findall(qn('w:tc'))): row.remove(cell)
        for j, value in enumerate(values):
            cell = deepcopy(src.rows[0 if i == 0 else 1].cells[0]._tc)
            row.append(cell)
        el.append(row)
        for j, (cell, value) in enumerate(zip(t.rows[-1].cells, values)):
            cell.width = Inches(widths[j])
            pr = cell._tc.get_or_add_tcPr()
            sh = pr.find(qn('w:shd')); sh.set(qn('w:fill'), '082A4A' if i==0 else ('E5EFF7' if i%2 else 'F5F8FA'))
            for child in list(cell._tc):
                if child.tag != qn('w:tcPr'): cell._tc.remove(child)
            par = cell.add_paragraph()
            par.paragraph_format.space_after = Pt(0)
            par.paragraph_format.line_spacing = 1.1
            font(par.add_run(value), 10, i==0, 'FFFFFF' if i==0 else '233447')
    for border in el.iter():
        if etree.QName(border).localname in ('top','bottom','left','right','insideH','insideV') and border.get(qn('w:color')):
            border.set(qn('w:color'), 'D9D9D9')
    p('').paragraph_format.space_after = Pt(0)
    return t

def link(label, url):
    par=p('', size=9)
    rel=doc.part.relate_to(url, 'http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink', is_external=True)
    h=OxmlElement('w:hyperlink'); h.set(qn('r:id'),rel)
    r=OxmlElement('w:r'); rp=OxmlElement('w:rPr'); c=OxmlElement('w:color');c.set(qn('w:val'),'23669B');rp.append(c);r.append(rp)
    tt=OxmlElement('w:t');tt.text=label;r.append(tt);h.append(r);par._p.append(h)
    par.paragraph_format.space_after=Pt(3)

# 封面沿用範本的雙標題、三欄狀態和四列資訊。
for i in range(8): p('', source=i)
for text, source in [('模式二馬達控制板',8),('使用與原理說明',9)]:
    par=p(text,source=source,size=30,bold=source==9)
    for r in par.runs:r.font.color.rgb=RGBColor(0,0,0)
picture(BOARD/'exports/pcb-3d.png',3.55)
for idx, texts in [(0,[['狀態\n硬體原型・未燒錄','','版本\nE 版','','更新日期\n2026 年 9 月 9 日']]),(1,[['板體','25 × 24 mm，雙面 PCB，四組螺絲端子'],['電源','2 粒 AA 串聯；配 Tamiya 980112M FA-130 馬達'],['操作','SW1 正轉 → SW2 反轉 → SW1 停止 → 鎖定 3 秒'],['內容','接線、操作邏輯、U1／U2、R／C、原理圖及銅線圖']])]:
    el=deepcopy(template.tables[idx]._tbl);body.insert(len(body)-1,el);t=Table(el,doc._body)
    for row,vals in zip(t.rows,texts):
        for cell,val in zip(row.cells,vals):
            first=cell.paragraphs[0];first.clear();font(first.add_run(val),9,color='233447')
            for extra in list(cell.paragraphs)[1:]:extra._p.getparent().remove(extra._p)
    if idx == 0:
        gap = p('')
        gap.paragraph_format.space_after = Pt(0)

heading('1  接線前先認識電路板')
p('E 版固定使用模式二，沒有 LED 和模式選擇跳線。四組綠色端子都用螺絲夾線；P5 的三個孔只供燒錄。以下以元件面朝上、S0 在上方為準。')
picture(BOARD/'exports/pcb-top.png',4.25,'圖 1　E 版元件正面。電池與馬達由右側進線；SW1、SW2 由下方進線。')
table(['接點與位置','第一條線','第二條線'],[
 ['P1 電池／右上','下方 1 腳 +：電池正極紅線','上方 2 腳 −：電池負極黑線'],
 ['P4 馬達／右側中下','下方 1 腳 M1／OUT1','上方 2 腳 M2／OUT2'],
 ['P2 SW1／左下','左側 1 腳：微動開關 NO','右側 2 腳：微動開關 COM'],
 ['P3 SW2／下方中間','左側 1 腳：微動開關 NO','右側 2 腳：微動開關 COM'],
 ],[1.45,2.825,2.825])
p('只用 2 粒 AA 串聯。不要接 3 粒 AA、USB 5 V 或 9 V。馬達兩端都由 U2 控制，不能把其中一端固定接 GND。',bold=True)
p('目前套件沒有韌體，也未製造或實測；空白 U1 不會執行本說明書的模式二流程。',size=10)

heading('2  安裝電線與操作順序')
sub('先接線再放入電池')
for s in [
 '1. 取出電池盒內全部電池，把 S0 撥至關機。核對 P1 的 +／−，四組端子不能互換用途。',
 '2. KF128-2.54 端子支援 AWG26–18 導線；電池及馬達建議用短 AWG22 銅線。剝去 5–6 mm 絕緣皮，鬆開螺絲後插線，鎖緊並輕拉確認。不要夾住絕緣皮或留下散開的銅絲。',
 '3. 帶線微動開關只用 COM 與 NO。用萬用表確認放開時不導通、按下時導通；NC 不接。開關兩根線的紅黑色不代表電源正負極。',
 '4. C4 是板外 100 nF 無極性陶瓷電容，焊在馬達兩個接點之間。接線後再檢查一次，才放入 2 粒 AA。',
]: p(s)
sub('以下是完成燒錄後應執行的模式二')
table(['步驟','動作','U1 狀態與馬達'],[
 ['① 開機','S0 接通電源','進入等待；IN1／IN2 = 00，馬達不受驅動'],
 ['② 起動','按下並放開 SW1','記住「正轉」狀態；輸出 10，放開仍繼續轉'],
 ['③ 返回','按下並放開 SW2','改為「反轉」狀態；輸出 01'],
 ['④ 停止','再次按下 SW1','輸出 00，馬達自由滑行停止'],
 ['⑤ 等候','保持停止 3 秒','忽略 SW1／SW2，之後回到等待狀態'],
 ],[.85,1.7,4.55])
p('沒有 LED 提示，操作時看馬達及機械裝置的動作。若方向與機構需要相反，先取出電池，再交換 P4 的兩條線。')
p('韌體尚待編寫：其他按鍵順序、同時按下、反轉前的過渡時間，以及堵轉和低電壓處理仍要定義。按住開關不應重複觸發，須加入去彈跳和按下事件判斷。',size=10)

heading('3  U1 如何決定下一步')
p('U1 是 ATtiny202 微控制器，負責讀取兩個開關、記住目前步驟和計算 3 秒。這些規則必須先寫成程式並燒錄進 U1；單靠焊好零件不會自動得到模式二。')
sub('一個按鍵訊號怎樣走到 U1')
p('SW1 放開時，U1 啟用的內置上拉電阻把 PA6 拉到高電位，讀作 1。SW1 壓下時，NO 與 COM 接通，PA6 經 P2 接到 GND，讀作 0。SW2 與 PA7 的讀法相同。')
p('因此「按下」是低電位，不是高電位。微動開關只傳送小電流訊號；馬達電流不經過微動開關或 U1。')
table(['U1 實體腳位','名稱','本板用途'],[
 ['1／8','VDD／GND','分別接 VBAT 供電與共同負極'],
 ['2','PA6','SW1 輸入，程式啟用內置上拉'],
 ['3','PA7','SW2 輸入，程式啟用內置上拉'],
 ['4','PA1','送到 U2 的 IN1'],
 ['5','PA2','送到 U2 的 IN2'],
 ['6','PA0／UPDI','經 P5 燒錄程式'],
 ['7','PA3／ADC','讀取 ISENSE 電壓；相關保護邏輯尚未實作'],
 ],[1.2,1.65,4.25])
sub('為甚麼放開 SW1 後馬達仍然轉')
p('U1 會保存「等待／正轉／反轉／鎖定」其中一個狀態。開關只用來改變狀態；放開開關不等於停止。到了反轉狀態，再按 SW1 才會進入停止和 3 秒鎖定。')
sub('去彈跳與計時不是由電容 C3 決定')
p('機械接點壓下的一瞬間可能快速通斷數次，程式要先確認訊號穩定，再當作一次按下。3 秒也是由 U1 的計時器和程式決定；更改 C1、C2 或 C3 不會直接改變這 3 秒。')
p('程式時鐘須符合供電電壓規格，本版以不高於 5 MHz 為目標。開機首先設定 IN1 = 0、IN2 = 0。',size=10)

heading('4  U2 如何令馬達正轉和反轉')
p('U2 是 DRV8213DSGR 馬達驅動器，內部有 H 橋，可以交換加在馬達兩端的電壓方向。U1 只送 IN1、IN2 兩個控制訊號，U2 才負責馬達電流。')
table(['IN1／IN2','OUT1／OUT2','馬達行為'],[
 ['0／0','兩端均高阻抗 High-Z','自由滑行；停止驅動，延遲後 U2 自動休眠'],
 ['1／0','OUT1 高、OUT2 低','定義為正轉，電流由 OUT1 經馬達流向 OUT2'],
 ['0／1','OUT1 低、OUT2 高','定義為反轉，電流方向相反'],
 ['1／1','兩端均低','電氣煞車；本模式二停止指令未採用此狀態'],
 ],[1.0,2.1,4.0])
p('表中的 1／0 表示邏輯高／低電位，不是 1 V／0 V。實際馬達轉向取決於 P4 接線及機械安裝；「正轉」不代表固定順時針。')
sub('正轉時的主要供電路徑')
p('電池正極 → Q1 → VBAT → U2 → OUT1 → 馬達 → OUT2 → U2 → GND → 電池負極',bold=True)
p('反轉時由 U2 交換 OUT1、OUT2 的高低電位。00 停止會讓馬達靠摩擦和負載慢慢停下，不能當作立即鎖住轉軸。')
sub('R1 把電流訊號轉成電壓並設定限流')
p('U2 的 IPROPI 輸出與馬達電流相關的小電流；R1 接在 IPROPI 與 GND 之間，把它變成 ISENSE 電壓，並讓 U2 設定限流門檻。R1 並沒有串在馬達主電流路徑上。')
p('GAINSEL 接 GND；依 TI 資料表，標稱比例為 205 µA/A，DSG 內部參考為 0.510 V。因此 R1 = 1.24 kΩ 對應的標稱限流約 2.006 A。',size=10)
p('限流是原型設定，並非整板可持續輸出 2 A 的保證，也不是堵轉自動停機。這顆 8 腳 DSG 版本沒有 16 腳 RTE 版本的獨立堵轉偵測功能。不要長時間卡住馬達測試。',bold=True)
link('資料依據　TI DRV8213 資料表：表 8-2 與電流感測章節','https://www.ti.com/lit/ds/symlink/drv8213.pdf')

heading('5  小電源開關與電阻值')
p('S0 是薄型滑動開關，只切換 Q1 的閘極訊號。Q1 是 P 通道 MOSFET，負責接通電池與整板之間的主要電流，所以 S0 可以做得很小。')
table(['部分','接法與數值','作用'],[
 ['Q1','DMP2035U-7；1=Gate、2=Source、3=Drain','Source 接 BATT+；Drain 輸出 VBAT'],
 ['S0 關機','1–2 接通','Gate 經 R2 接近 Source 電位，Q1 關斷'],
 ['S0 開機','2–3 接通','Gate 被拉向 GND，Q1 導通'],
 ['R2','1 kΩ，1%，0603','限制閘極切換瞬間電流，並非馬達限流'],
 ['R3','100 kΩ，1%，0603','把 Gate 拉回 Source，開關換檔間隙預設關機'],
 ],[1.0,2.7,3.4])
sub('電阻先看用途才選數值')
p('R1 按 U2 的原廠限流公式選值；R2 依開關可承受的瞬間閘極電流選值；R3 要能可靠拉回閘極，同時避免持續浪費太多電。它們不能因為外觀相同就互換。')
p('R1 = 0.510 V ÷（205 µA/A × 2 A）≈ 1.244 kΩ，採用標準值 1.24 kΩ、1%。IC 與電阻誤差會令實際門檻偏移。')
p('R2 = 1 kΩ 時，以 3.3 V 估算的初始電流約 3.3 V ÷ 1,000 Ω = 3.3 mA，低於指定 S0 的 50 mA 額定。')
p('開機後 R2、R3 形成 101 kΩ 支路；3 V 下電流約 29.7 µA，Q1 的 Gate 比 Source 低約 2.97 V。')
sub('關掉 S0 不等於完全隔離所有電源')
p('Q1 不是反接保護，也不是雙向隔離開關。若燒錄器由 P5 供電，電流可經 Q1 的本體二極體回流到電池端，必須先取出電池盒內全部電池，不能只關 S0。',bold=True)
link('Q1 資料表　Diodes DMP2035U','https://www.diodes.com/datasheet/download/DMP2035U.pdf')

heading('6  電容與讀圖符號')
table(['元件','數值與位置','為甚麼需要'],[
 ['C1','100 nF／16 V X7R，靠近 U1','供應晶片短促的電流需求，減少電源高頻干擾'],
 ['C2','100 nF／16 V X7R，靠近 U2','在驅動器旁作局部去耦，連線要短'],
 ['C3','100 µF／6.3 V，有極性，板上','儲備電荷，減少馬達啟動時的供電波動；容量仍需實測'],
 ['C4','100 nF，無極性，耐壓至少 10 V，板外','直接焊在馬達兩端，抑制電刷產生的雜訊'],
 ],[.65,2.55,3.9])
p('C3 正極接 VBAT，負極接 GND，裝反可能損壞。C1、C2、C4 是無極性陶瓷電容。100 nF = 0.1 µF；100 µF 的容量是它的 1,000 倍。')
sub('容量並非越大越好')
p('先用 IC 原廠建議的去耦值，再看馬達、電池、導線及容許壓降。粗略估算可用 C ≈ I × Δt ÷ ΔV，但它只估計電容短暫供電的部分，未包括電池與導線阻抗。')
p('例如額外 0.2 A 持續 0.1 ms，若容許下降 0.2 V，粗算需要 100 µF。這只是教學例子，不是本板量測結果；C3 的 100 µF 仍是原型初值。')
sub('原理圖上的名字怎樣看')
table(['記號','讀法','在本板的例子'],[
 ['U／Q／R／C／P／S','IC／電晶體／電阻／電容／接點／開關','U1 是位號；旁邊的 2 才是實體腳位編號'],
 ['BATT+／VBAT','開關前／開關後的正電源','Q1 把 BATT+ 接通到 VBAT'],
 ['GND','共同參考電位和回流路徑','接電池負極，並非建築物的保護接地'],
 ['同名標籤','同一個電氣網絡，不必畫長線連起','U1 的 SW1 與 P2 的 SW1 相連'],
 ['交點圓點／#FLG','圓點代表相連；#FLG 是檢查標記','線條交叉不能只憑接近判定；#FLG 不需購買'],
 ],[1.25,2.95,2.9])

heading('7  完整原理圖')
picture(BOARD/'exports/schematic-preview.png',7.1,'圖 2　由 E 版 KiCad 原理圖直接輸出。附圖包另含可放大的原始向量 PDF。')
sub('依照電源、訊號、輸出的次序閱讀')
p('上方：P1 → S0／R2／R3／Q1 是供電控制；C1、C2、C3 接在 VBAT 與 GND 之間。')
p('中間：P2／P3 把開關訊號送入 U1；U1 的 IN1／IN2 送至 U2；U2 由 P4 接馬達。ISENSE 把電流感測電壓送回 U1。')
p('下方：S1、S2 是板外微動開關；M1 是板外馬達；C4 也在馬達上。P5 是燒錄孔。原理圖上的相對位置不等於 PCB 上的實際位置。')
p('讀細小腳位時，請放大電子文件，或開啟附圖包內的「完整原理圖.pdf」。圖面是實際電路資料，不是實物拍攝照片。',size=10)

heading('8  PCB 正面銅線佈局')
picture(ROOT/'assets/pcb-front-copper.png',6.65,'圖 3　F.Cu 正面銅層與正面絲印。元件面朝向讀者，S0 在上方。')
p('紅色都是銅，包括狹長走線、焊盤及大片 GND 鋪銅；白色間隙是銅之間的隔離區，不是電線。淺黃色是元件名稱與外框絲印，不導電。')
p('有孔的焊盤讓端子焊腳穿過；較小的鍍通孔可把正、反面銅層相連。不是每個圓孔都屬同一網絡，必須對照原理圖或 KiCad 網絡高亮。')
p('實際板體 25 × 24 mm；圖片已放大，不能量圖決定尺寸。所有元件裝在正面，綠色防焊層會覆蓋大部分銅線。',size=10)

heading('9  PCB 背面銅線佈局')
picture(ROOT/'assets/pcb-back-copper.png',6.65,'圖 4　B.Cu 背面銅層與背面絲印，已鏡像成從背面觀看的方向。')
p('藍色都是背面銅，白色為隔離區。圖中的電池端子在左上，是因為把板翻面；正面接線圖中的電池仍在右上。不要把兩種視角的位置直接混用。')
p('較寬的線用於電池與馬達電流，較細的線傳送控制訊號。大片 GND 鋪銅提供回流與散熱；走線只出現在其中一面時，可能經鍍通孔換到另一面。')
p('本板是雙面 FR-4、厚 1.6 mm、銅厚 1 oz。這些 KiCad 圖是佈局預覽；製造仍須使用 E 版的 Gerber 與鑽孔檔。',size=10)

heading('10  交給賣家燒錄與首次驗證')
p('目前只有硬體設計，沒有可交給賣家燒錄的 HEX 檔。先完成模式二韌體與測試，再把程式檔、晶片設定及操作驗收要求一起交給賣家；PCBA 焊接服務不一定包括代燒錄。')
table(['項目','需要交代的內容','目前狀態'],[
 ['燒錄晶片','ATtiny202-SSNR，使用 UPDI；U2 不需燒程式','U1 韌體尚未提供'],
 ['P5 腳位','正面左至右：1 VBAT、2 UPDI、3 GND','三個 2.54 mm 間距孔，可不裝排針'],
 ['供電方式','燒錄器供電時取出全部 AA；只選一個供電來源','不能只關 S0'],
 ['功能驗收','開機停止；SW1→正轉、SW2→反轉、SW1→停止；鎖定 3 秒','須燒錄後實測'],
 ],[1.0,4.45,1.65])
sub('第一次上電要確認甚麼')
p('先檢查電池極性、C3 方向、端子夾線與焊接短路。U2 是底部接點 WSON，中央散熱焊盤也必須焊好，宜由有相應工藝的裝配廠處理。')
p('燒錄後先作空載操作，再接實際機構測試啟動、反轉、停止及溫升。若異常發熱或馬達卡住，立即關機並取出電池。約 2 A 限流不能代替堵轉停機策略。')
p('E 版記錄：已啟用的 ERC／DRC 規則均為 0 問題，未接線與原理圖一致性問題為 0。這是設計檢查，不能證明實際電源、散熱或程式已正常。')
p('3D 預覽使用通用及尺寸外形模型；S0、U2 和端子細節為示意，未用於精密外殼干涉驗證。本板尚未製造、焊接、燒錄或實機測試。',size=10)
sub('資料來源與圖面')
for label,url in [
 ('Microchip ATtiny202 資料表與腳位','https://ww1.microchip.com/downloads/aemDocuments/documents/MCU08/ProductDocuments/DataSheets/ATtiny202-204-402-404-406-DataSheet-DS40002318A.pdf'),
 ('TI DRV8213 資料表與 H 橋邏輯','https://www.ti.com/lit/ds/symlink/drv8213.pdf'),
 ('KEFA KF128-2.54 原廠端子圖面','https://datasheet.lcsc.com/datasheet/pdf/7193d1712e806c94a94291057681f262.pdf?productCode=C474920'),
 ('SHOU HAN MSK12C02 原廠開關圖面','https://datasheet.lcsc.com/datasheet/pdf/5162155576bfd231c35aa9a893d25c8c.pdf?productCode=C431540'),
 ('Panasonic C3 規格　EEEFK0J101UR','https://industrial.panasonic.com/ww/products/pt/aluminum-cap-smd/models/EEEFK0J101UR'),
 ('Tamiya 980112M 馬達規格','https://www.pololu.com/product/77/specs'),
]:link(label,url)

for sec in doc.sections:
    for par in sec.footer.paragraphs:
        if 'Organization Name' in par.text:
            par.clear();font(par.add_run('模式二馬達控制板  |  E 版使用與原理說明'),9,color='5B7085')
            par.add_run('  |  ')
            field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');par._p.append(field)

doc.core_properties.title='模式二馬達控制板 E 版使用與原理說明'
doc.core_properties.subject='接線、模式二邏輯、元件作用與 KiCad 圖面'
doc.core_properties.author=''
doc.core_properties.last_modified_by=''
temp=ROOT/'qa/manual-working.docx';doc.save(temp)
# 只換內容零件，保留原範本其餘套件資料，避免 python-docx 改寫無關 XML。
editable={'word/document.xml','word/_rels/document.xml.rels','[Content_Types].xml','docProps/core.xml','word/footer1.xml'}
with ZipFile(REF) as original, ZipFile(temp) as changed, ZipFile(OUT,'w',ZIP_DEFLATED) as final:
    for name in changed.namelist():
        final.writestr(name,changed.read(name) if name in editable or name not in original.namelist() else original.read(name))
    for name in original.namelist():
        if name not in changed.namelist():final.writestr(name,original.read(name))

with ZipFile(REF) as a, ZipFile(OUT) as b:
    unchanged=[n for n in a.namelist() if n not in editable]
    assert all(a.read(n)==b.read(n) for n in unchanged)
    assert len(Document(OUT).sections)==1
    assert Document(OUT).sections[0]._sectPr.xml==template.sections[0]._sectPr.xml
    assert len(Document(OUT).inline_shapes)==5
net=etree.parse(str(BOARD/'exports/mode2-netlist.xml'))
actual={}
for netel in net.findall('.//nets/net'):
    for node in netel.findall('node'):actual[(node.get('ref'),node.get('pin'))]=netel.get('name')
for key,name in {('U1','2'):'SW1',('U1','3'):'SW2',('U1','4'):'IN1',('U1','5'):'IN2',('P1','1'):'BATT+',('P1','2'):'GND',('P5','1'):'VBAT'}.items():
    assert actual[key].lstrip('/')==name,(key,actual[key])
(ROOT/'qa/validation.json').write_text(json.dumps({'preserved_template_parts':len(unchanged),'sections':1,'images':5,'pinmap_check':'pass','source_sha256':{n:sha256((BOARD/n).read_bytes()).hexdigest() for n in ['mode2-motor.kicad_sch','mode2-motor.kicad_pcb']},'reference_sha256':sha256(REF.read_bytes()).hexdigest()},ensure_ascii=False,indent=2),encoding='utf-8')
print(OUT)
